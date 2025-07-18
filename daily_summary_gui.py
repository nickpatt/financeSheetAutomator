#!/usr/bin/env python3
"""
Daily Summary Generator GUI
A Windows GUI version of the daily summary generator
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import threading
import os
import sys
import subprocess

# Import the main functionality from the original script
from daily_summary_generator import generate_summary, scan_available_project_files

class DailySummaryGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Daily Summary Generator")
        self.root.geometry("600x500")
        self.root.resizable(True, True)
        
        # Variables
        self.target_date_var = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.output_dir_var = tk.StringVar(value="reports")
        self.year_vars = {}  # Will store year checkboxes
        self.available_years = []  # Will store available years
        
        # Create GUI
        self.create_widgets()
        
    def create_widgets(self):
        # Main frame
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # Configure grid weights
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        main_frame.columnconfigure(1, weight=1)
        
        # Title
        title_label = ttk.Label(main_frame, text="Daily Summary Generator", 
                               font=('Arial', 16, 'bold'))
        title_label.grid(row=0, column=0, columnspan=3, pady=(0, 20))
        
        # Date selection
        ttk.Label(main_frame, text="Target Date:").grid(row=1, column=0, sticky=tk.W, pady=5)
        date_frame = ttk.Frame(main_frame)
        date_frame.grid(row=1, column=1, sticky=(tk.W, tk.E), pady=5, padx=(10, 0))
        
        self.date_entry = ttk.Entry(date_frame, textvariable=self.target_date_var, width=15)
        self.date_entry.pack(side=tk.LEFT)
        
        ttk.Button(date_frame, text="Today", command=self.set_today,
                  width=8).pack(side=tk.LEFT, padx=(10, 0))
        
        # Date format help
        ttk.Label(main_frame, text="(Format: YYYY-MM-DD)", 
                 font=('Arial', 8), foreground='gray').grid(row=2, column=1, sticky=tk.W, padx=(10, 0))
        
        # Output directory
        ttk.Label(main_frame, text="Output Directory:").grid(row=3, column=0, sticky=tk.W, pady=(15, 5))
        
        output_frame = ttk.Frame(main_frame)
        output_frame.grid(row=3, column=1, sticky=(tk.W, tk.E), pady=(15, 5), padx=(10, 0))
        output_frame.columnconfigure(0, weight=1)
        
        self.output_entry = ttk.Entry(output_frame, textvariable=self.output_dir_var)
        self.output_entry.grid(row=0, column=0, sticky=(tk.W, tk.E))
        
        ttk.Button(output_frame, text="Browse", command=self.browse_output_dir,
                  width=8).grid(row=0, column=1, padx=(10, 0))
        
        # Year selection
        year_frame = ttk.LabelFrame(main_frame, text="Select Years to Process", padding="10")
        year_frame.grid(row=4, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(15, 0))
        
        # Scan for available files and create checkboxes
        self.setup_year_selection(year_frame)
        
        # Configuration info
        config_frame = ttk.LabelFrame(main_frame, text="Configuration", padding="10")
        config_frame.grid(row=5, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(20, 0))
        config_frame.columnconfigure(1, weight=1)
        
        ttk.Label(config_frame, text="Primary Data Directory:").grid(row=0, column=0, sticky=tk.W)
        ttk.Label(config_frame, text="N:\\Project List\\", 
                 font=('Arial', 9, 'bold')).grid(row=0, column=1, sticky=tk.W, padx=(10, 0))
        
        ttk.Label(config_frame, text="Fallback Directories:").grid(row=1, column=0, sticky=tk.W)
        ttk.Label(config_frame, text="quarterly sheets, reports", 
                 font=('Arial', 9, 'bold')).grid(row=1, column=1, sticky=tk.W, padx=(10, 0))
        
        ttk.Label(config_frame, text="Supported Formats:").grid(row=2, column=0, sticky=tk.W)
        ttk.Label(config_frame, text=".xlsx, .xlsm", 
                 font=('Arial', 9, 'bold')).grid(row=2, column=1, sticky=tk.W, padx=(10, 0))
        
        # Progress bar
        self.progress_var = tk.StringVar(value="Ready")
        ttk.Label(main_frame, text="Status:").grid(row=6, column=0, sticky=tk.W, pady=(20, 5))
        self.status_label = ttk.Label(main_frame, textvariable=self.progress_var, 
                                     font=('Arial', 9), foreground='blue')
        self.status_label.grid(row=6, column=1, sticky=tk.W, pady=(20, 5), padx=(10, 0))
        
        # Progress bar
        self.progress_bar = ttk.Progressbar(main_frame, mode='indeterminate')
        self.progress_bar.grid(row=7, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(5, 20))
        
        # Buttons
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=8, column=0, columnspan=3, pady=(10, 0))
        
        self.generate_button = ttk.Button(button_frame, text="Generate Summary", 
                                         command=self.start_generation, width=20)
        self.generate_button.pack(side=tk.LEFT, padx=(0, 10))
        
        ttk.Button(button_frame, text="Exit", command=self.root.quit, 
                  width=12).pack(side=tk.LEFT)
        
        # Output text area
        output_frame = ttk.LabelFrame(main_frame, text="Output Log", padding="5")
        output_frame.grid(row=9, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(20, 0))
        output_frame.columnconfigure(0, weight=1)
        output_frame.rowconfigure(0, weight=1)
        main_frame.rowconfigure(9, weight=1)
        
        # Text widget with scrollbar
        text_frame = ttk.Frame(output_frame)
        text_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        text_frame.columnconfigure(0, weight=1)
        text_frame.rowconfigure(0, weight=1)
        
        self.output_text = tk.Text(text_frame, height=8, wrap=tk.WORD)
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.output_text.yview)
        self.output_text.configure(yscrollcommand=scrollbar.set)
        
        self.output_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        
    def set_today(self):
        """Set date to today"""
        self.target_date_var.set(datetime.now().strftime("%Y-%m-%d"))
        
    def browse_output_dir(self):
        """Browse for output directory"""
        directory = filedialog.askdirectory(initialdir=self.output_dir_var.get())
        if directory:
            self.output_dir_var.set(directory)
    
    def setup_year_selection(self, parent_frame):
        """Set up year selection checkboxes"""
        # Scan for available files
        try:
            available_files = scan_available_project_files(2023, 2030)
            self.available_years = [year for year, _ in available_files]
            
            if not self.available_years:
                ttk.Label(parent_frame, text="No Project List files found in any location.", 
                         foreground='red').pack(pady=10)
                return
                
            # Create instruction label
            ttk.Label(parent_frame, text="Select which years to include in the summary:").pack(anchor=tk.W, pady=(0, 10))
            
            # Create checkboxes in a scrollable frame if needed
            checkbox_frame = ttk.Frame(parent_frame)
            checkbox_frame.pack(fill=tk.X)
            
            # Create checkboxes in rows of 4
            row = 0
            col = 0
            for year in self.available_years:
                # Default to checked for years 2023-2025, unchecked for others
                default_checked = year in ['2023', '2024', '2025']
                var = tk.BooleanVar(value=default_checked)
                self.year_vars[year] = var
                
                checkbox = ttk.Checkbutton(checkbox_frame, text=year, variable=var)
                checkbox.grid(row=row, column=col, sticky=tk.W, padx=(0, 15), pady=2)
                
                col += 1
                if col >= 4:  # 4 columns per row
                    col = 0
                    row += 1
                    
            # Add Select All / Deselect All buttons
            button_frame = ttk.Frame(parent_frame)
            button_frame.pack(fill=tk.X, pady=(10, 0))
            
            ttk.Button(button_frame, text="Select All", 
                      command=self.select_all_years, width=12).pack(side=tk.LEFT)
            ttk.Button(button_frame, text="Deselect All", 
                      command=self.deselect_all_years, width=12).pack(side=tk.LEFT, padx=(10, 0))
            
            # Show count of available files
            count_label = ttk.Label(parent_frame, 
                                   text=f"Found {len(self.available_years)} available Project List files", 
                                   font=('Arial', 8), foreground='green')
            count_label.pack(anchor=tk.W, pady=(10, 0))
            
        except Exception as e:
            error_label = ttk.Label(parent_frame, text=f"Error scanning files: {str(e)}", 
                                   foreground='red')
            error_label.pack(pady=10)
    
    def select_all_years(self):
        """Select all available years"""
        for var in self.year_vars.values():
            var.set(True)
    
    def deselect_all_years(self):
        """Deselect all years"""
        for var in self.year_vars.values():
            var.set(False)
    
    def get_selected_years(self):
        """Get list of selected years"""
        return [year for year, var in self.year_vars.items() if var.get()]
            
    def log_message(self, message):
        """Add message to output log"""
        self.output_text.insert(tk.END, message + "\n")
        self.output_text.see(tk.END)
        self.root.update_idletasks()
        
    def validate_inputs(self):
        """Validate user inputs"""
        # Validate date
        try:
            datetime.strptime(self.target_date_var.get(), "%Y-%m-%d")
        except ValueError:
            messagebox.showerror("Invalid Date", "Please enter a valid date in YYYY-MM-DD format")
            return False
            
        # Validate output directory
        output_dir = self.output_dir_var.get().strip()
        if not output_dir:
            messagebox.showerror("Invalid Directory", "Please specify an output directory")
            return False
        
        # Validate that at least one year is selected
        selected_years = self.get_selected_years()
        if not selected_years:
            messagebox.showerror("No Years Selected", "Please select at least one year to process")
            return False
            
        return True
        
    def start_generation(self):
        """Start the summary generation in a separate thread"""
        if not self.validate_inputs():
            return
            
        # Disable button and start progress
        self.generate_button.config(state='disabled')
        self.progress_bar.start(10)
        self.progress_var.set("Generating summary...")
        self.output_text.delete(1.0, tk.END)
        
        # Start generation in thread
        thread = threading.Thread(target=self.run_generation)
        thread.daemon = True
        thread.start()
        
    def run_generation(self):
        """Run the actual summary generation"""
        try:
            target_date = datetime.strptime(self.target_date_var.get(), "%Y-%m-%d").date()
            output_dir = self.output_dir_var.get().strip()
            selected_years = self.get_selected_years()
            
            self.log_message(f"Starting generation for {target_date}...")
            self.log_message(f"Output directory: {output_dir}")
            self.log_message(f"Selected years: {', '.join(selected_years)}")
            self.log_message("-" * 50)
            
            # Redirect stdout to capture print statements
            import io
            import contextlib
            
            old_stdout = sys.stdout
            captured_output = io.StringIO()
            
            with contextlib.redirect_stdout(captured_output):
                success = generate_summary(target_date, output_dir, selected_years)
            
            # Restore stdout
            sys.stdout = old_stdout
            
            # Display captured output
            output_lines = captured_output.getvalue().split('\n')
            for line in output_lines:
                if line.strip():
                    self.root.after(0, self.log_message, line)
            
            # Show result
            if success:
                self.root.after(0, self.on_success, target_date, output_dir)
            else:
                self.root.after(0, self.on_error, "Generation failed. Check the output log for details.")
                
        except Exception as e:
            self.root.after(0, self.on_error, f"Error: {str(e)}")
            
    def on_success(self, target_date, output_dir):
        """Handle successful generation"""
        self.progress_bar.stop()
        self.progress_var.set("Generation completed successfully!")
        self.generate_button.config(state='normal')
        
        # Show success message with options
        result = messagebox.askyesno("Success", 
                                   f"Summary generated successfully!\n\n"
                                   f"Files saved to: {output_dir}\n\n"
                                   f"Would you like to open the output folder?")
        if result:
            try:
                # Open the output directory
                if os.name == 'nt':  # Windows
                    os.startfile(output_dir)
                else:  # Other OS
                    subprocess.run(['open' if sys.platform == 'darwin' else 'xdg-open', output_dir])
            except Exception as e:
                messagebox.showwarning("Cannot Open Folder", f"Could not open folder: {e}")
                
    def on_error(self, error_message):
        """Handle generation error"""
        self.progress_bar.stop()
        self.progress_var.set("Generation failed")
        self.generate_button.config(state='normal')
        self.log_message(f"ERROR: {error_message}")
        messagebox.showerror("Error", error_message)

def check_dependencies():
    """Check if required packages are installed"""
    required_packages = ['pandas', 'docx', 'openpyxl']
    missing_packages = []
    
    for package in required_packages:
        try:
            if package == 'docx':
                import docx
            else:
                __import__(package)
        except ImportError:
            missing_packages.append(package if package != 'docx' else 'python-docx')
    
    if missing_packages:
        messagebox.showerror("Missing Dependencies", 
                           f"The following packages are required but not installed:\n\n"
                           f"{', '.join(missing_packages)}\n\n"
                           f"Please install them using:\n"
                           f"pip install {' '.join(missing_packages)}")
        return False
    return True

def main():
    """Main function"""
    # Check dependencies first
    if not check_dependencies():
        return
        
    # Create and run the GUI
    root = tk.Tk()
    app = DailySummaryGUI(root)
    
    # Center window on screen
    root.update_idletasks()
    width = root.winfo_width()
    height = root.winfo_height()
    x = (root.winfo_screenwidth() // 2) - (width // 2)
    y = (root.winfo_screenheight() // 2) - (height // 2)
    root.geometry(f'{width}x{height}+{x}+{y}')
    
    root.mainloop()

if __name__ == "__main__":
    main() 